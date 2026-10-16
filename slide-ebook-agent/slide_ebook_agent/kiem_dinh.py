"""Tác tử kiểm định: đồng bộ ebook ↔ slide, chữ tràn, cỡ chữ, link thương hiệu, trang thưa, contact sheet -> bao-cao-qa.md."""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

import pypdfium2 as pdfium
from PIL import Image
from pypdf import PdfReader

from thuong_hieu import brand

BODY_ROLES = {"body", "compare-item", "side", "summary", "agenda"}


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\*\*|\[\[#\w+\||\]\]", "", s)).strip().lower()


def pdf_texts(path: Path) -> list[str]:
    d = pdfium.PdfDocument(str(path))
    out = [re.sub(r"\s+", " ", d[i].get_textpage().get_text_range()).lower() for i in range(len(d))]
    d.close()
    return out


def pdf_links(path: Path) -> Counter:
    c = Counter()
    for pg in PdfReader(str(path)).pages:
        for a in pg.get("/Annots") or []:
            u = a.get_object().get("/A", {}).get("/URI")
            if u:
                c[str(u).rstrip("/")] += 1
    return c


def pptx_links(path: Path) -> Counter:
    from pptx import Presentation

    c = Counter()
    for sl in Presentation(str(path)).slides:
        for sh in sl.shapes:
            try:
                if sh.click_action.hyperlink.address:
                    c[sh.click_action.hyperlink.address.rstrip("/")] += 1
            except Exception:
                pass
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.hyperlink.address:
                            c[r.hyperlink.address.rstrip("/")] += 1
    return c


def contact_sheet(pdf: Path, out_prefix: Path, per_row: int, per_sheet: int, scale: float) -> list[Path]:
    d = pdfium.PdfDocument(str(pdf))
    ims = [d[i].render(scale=scale).to_pil() for i in range(len(d))]
    d.close()
    w, h = ims[0].size
    files = []
    for s in range(0, len(ims), per_sheet):
        chunk = ims[s:s + per_sheet]
        rows = -(-len(chunk) // per_row)
        sheet = Image.new("RGB", (per_row * w + (per_row + 1) * 6, rows * h + (rows + 1) * 6), "#9aa3b5")
        for i, im in enumerate(chunk):
            sheet.paste(im.resize((w, h)), (6 + (i % per_row) * (w + 6), 6 + (i // per_row) * (h + 6)))
        f = out_prefix.parent / f"{out_prefix.name}_{s // per_sheet + 1:02d}.jpg"
        sheet.save(f, quality=78)
        files.append(f)
    return files


def check(doc: dict, deck, overflow: list[dict], ebook_pdf: Path | None, slide_pdf: Path | None,
          slide_pptx: Path | None, ebook_info: dict | None, qa_dir: Path) -> tuple[str, int]:
    qa_dir.mkdir(parents=True, exist_ok=True)
    b = brand()
    errors, warns, lines = [], [], []

    # 1. Đồng bộ chương
    eb = pdf_texts(ebook_pdf) if ebook_pdf else []
    lines += ["## Đồng bộ ebook ↔ slide", "", "| # | Chương | Trang ebook | Slide |", "|---|---|---|---|"]
    for i, ch in enumerate(doc["chapters"], 1):
        key = norm(ch["title"])[:60]
        flat = [t.replace(" ", "") for t in eb]   # PDF giữ khoảng giãn chữ: "C H Ư Ơ N G 0 1"
        page = next((str(p + 1) for p, t in enumerate(flat) if key.replace(" ", "") in t and f"chương{i:02d}" in t
                     and "mụclục" not in t), "–") if eb else "(không dựng)"
        sl = [str(k + 1) for k, s in enumerate(deck.slides) if s["chapter"] == ch["id"]] if deck else []
        slide_txt = f"{sl[0]}–{sl[-1]} ({len(sl)})" if sl else ("–" if deck else "(không dựng)")
        lines.append(f"| {i} | {ch['title']} | {page} | {slide_txt} |")
        if eb and page == "–":
            errors.append(f"Ebook thiếu chương {i} «{ch['title']}».")
        if deck and len(sl) < 2:
            errors.append(f"Slide thiếu nội dung cho chương {i} «{ch['title']}».")
    lines.append("")

    # 2. Slide: tràn chữ, cỡ chữ
    if deck:
        for o in overflow:
            errors.append(f"Slide {o['slide'] + 1} vẫn tràn chữ sau khi rút gọn: {'; '.join(o['problems'])}")
        small = Counter()
        for k, s in enumerate(deck.slides):
            for e in s["elements"]:
                if e["kind"] == "text" and e["role"] in BODY_ROLES and e["size"] < 14:
                    small[k + 1] += 1
        if small:
            warns.append("Chữ nội dung dưới 14pt ở slide: " + ", ".join(map(str, sorted(small))) + " (đọc trên máy chiếu có thể khó).")
        layouts = Counter(s["kind"] for s in deck.slides)
        lines += ["## Bộ slide", "", f"- {len(deck.slides)} slide; bố cục: " + ", ".join(f"{k} ×{v}" for k, v in layouts.most_common()), ""]

    # 3. Link thương hiệu
    need = {b["site"]: "ai.slim.vn", **{e["link"]: e["name"] for e in b["ecosystem"]}, **{s["link"]: s["name"] for s in b["social"]}}
    if doc.get("course_enabled", True):
        need[b["course"]["url"]] = "Khoá học"
    for label, links in [("Ebook PDF", pdf_links(ebook_pdf) if ebook_pdf else None),
                         ("Slide PDF", pdf_links(slide_pdf) if slide_pdf else None),
                         ("Slide PPTX", pptx_links(slide_pptx) if slide_pptx else None)]:
        if links is None:
            continue
        miss = [n for u, n in need.items() if u.rstrip("/") not in links]
        bad = [u for u in links if "/login" in u]
        lines.append(f"- **{label}**: {sum(links.values())} link" + (f" — thiếu: {', '.join(miss)}" if miss else " — đủ link thương hiệu"))
        if miss:
            errors.append(f"{label} thiếu link: {', '.join(miss)}")
        if bad:
            errors.append(f"{label} có link sai: {', '.join(bad)}")
    lines.append("")

    # 4. Ebook: mục lục, trang thưa
    if ebook_info:
        for m in ebook_info.get("toc_missing", []):
            warns.append(f"Mục lục ebook: {m.strip()}")
    if eb:
        sparse = [i + 1 for i, t in enumerate(eb) if len(t.strip()) < 120]
        if sparse:
            warns.append(f"Trang ebook ít chữ (kiểm tra trang trắng thừa): {sparse}")
        lines.append(f"- Ebook: {len(eb)} trang")

    # 5. Contact sheet
    sheets = []
    if ebook_pdf:
        sheets += contact_sheet(ebook_pdf, qa_dir / "ebook", 5, 10, 0.42)
    if slide_pdf:
        sheets += contact_sheet(slide_pdf, qa_dir / "slide", 3, 12, 0.6)

    head = ["# Báo cáo kiểm định", "", f"**{doc['title']}**", ""]
    if errors:
        head += ["## ❌ Lỗi cần sửa", *[f"- {e}" for e in errors], ""]
    else:
        head += ["## ✅ Không có lỗi chặn", ""]
    if warns or doc.get("warnings"):
        head += ["## ⚠ Cần xem", *[f"- {w}" for w in warns], *[f"- (biên tập) {w}" for w in doc.get("warnings", [])], ""]
    tail = ["## Ảnh kiểm tra từng trang", *[f"- {s.name}" for s in sheets], "",
            "> Mở từng contact sheet để xem bằng mắt: ảnh cắt dở, khoảng trống lớn, chữ sát mép."]
    return "\n".join(head + lines + tail), len(errors)
