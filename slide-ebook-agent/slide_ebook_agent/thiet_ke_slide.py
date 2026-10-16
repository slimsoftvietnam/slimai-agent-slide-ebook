"""Tác tử thiết kế slide: bản thảo -> danh sách phần tử có toạ độ -> đo chữ trong Chromium -> PDF + PPTX gốc (sửa được).

Một danh sách phần tử duy nhất cho cả hai đầu ra => PPTX và PDF cùng toạ độ, cùng cỡ chữ.
Đơn vị: point. Slide 16:9 = 720 x 405 pt (10 x 5.625 inch).
"""
from __future__ import annotations

import copy
import html
import re
from pathlib import Path

from PIL import Image

from thuong_hieu import brand, hex_rgb, logo, qr_png, social_pngs

W, H = 720, 405
M = 40                              # lề trái/phải
FONT, MONO = "Arial", "Courier New"  # có sẵn trên mọi máy, đủ dấu tiếng Việt
LH = 1.22                           # line-height dùng chung cho HTML và PPTX
NAVY, INK, MUTED, LINE, SOFT = "#0b2a5b", "#1d2433", "#5b6475", "#e3e7ef", "#f4f7fc"
DARK, DARK2 = "#0b1a3a", "#10285a"
PINK, PURPLE, CYAN, GREEN, YELLOW = "#e8318a", "#8a4dea", "#12b5e5", "#3cc26b", "#f6c21c"
GRAD = [PINK, PURPLE, CYAN]
ACCENTS = [PINK, CYAN, GREEN, PURPLE, YELLOW]


# ---------------------------------------------------------------- phần tử
def rect(x, y, w, h, fill=None, grad=None, radius=0, alpha=1.0, line=None, name="", line_w=0.75, shadow=False, rot=0, grad_angle=45):
    return dict(kind="rect", x=x, y=y, w=w, h=h, fill=fill, grad=grad, radius=radius, alpha=alpha, line=line, name=name,
                line_w=line_w, shadow=shadow, rot=rot, grad_angle=grad_angle)


def oval(x, y, w, h, fill=None, line=None, line_w=0.75, alpha=1.0, name=""):
    return dict(kind="oval", x=x, y=y, w=w, h=h, fill=fill, line=line, line_w=line_w, alpha=alpha, name=name)


def line(x1, y1, x2, y2, color, width=0.75, name=""):
    """Đường thẳng bất kỳ = hình chữ nhật mảnh xoay quanh tâm (HTML và PPTX xoay giống nhau)."""
    import math
    L = math.hypot(x2 - x1, y2 - y1)
    cx, cy = (x1 + x2) / 2, (y1 + y2) / 2
    return rect(cx - L / 2, cy - width / 2, L, width, fill=color, rot=math.degrees(math.atan2(y2 - y1, x2 - x1)), name=name)


def text(x, y, w, h, paras, size, role, *, color=INK, bold=False, italic=False, align="left", valign="top",
         min_size=None, bullet=None, space=0, font=FONT, link=None, caps=False, spacing_chars=0, name="",
         accent=None, accent_font=None, weight=None, accent_italic=True):
    if isinstance(paras, str):
        paras = [paras]
    return dict(kind="text", x=x, y=y, w=w, h=h, paras=[p for p in paras if p], size=size, role=role, color=color,
                bold=bold, italic=italic, align=align, valign=valign, min=min_size or size, bullet=bullet, space=space,
                font=font, link=link, caps=caps, cs=spacing_chars, name=name, accent=accent, accent_font=accent_font, weight=weight,
                accent_italic=accent_italic)


def image(x, y, w, h, path, link=None, fit="cover", name="", rot=0):
    return dict(kind="image", x=x, y=y, w=w, h=h, path=str(path), link=link, fit=fit, name=name, rot=rot)


# ---------------------------------------------------------------- khung chung
class Deck:
    sans = FONT                 # font nội dung
    font_links: list = []       # Google Fonts cho HTML/PDF (PPTX dùng font cài trên máy)
    transition = None           # "morph" -> PPTX dùng Morph, HTML dùng View Transitions

    def __init__(self, doc: dict, work: Path, images: dict[str, Path], social: dict[str, Path]):
        self.doc, self.work, self.images, self.social = doc, work, images, social
        self.b = brand()
        self.slides: list[dict] = []   # {elements, bg, notes, kind, chapter, footer}
        self.crop_dir = work / "slide-anh"
        self.crop_dir.mkdir(exist_ok=True)

    # ảnh: cắt sẵn đúng tỉ lệ khung để HTML và PPTX hiển thị y hệt
    def img(self, image_id, x, y, w, h, link=None, name=""):
        src = self.images.get(image_id or "")
        if not src:
            return []
        out = self.crop_dir / f"{Path(src).stem}_{int(w)}x{int(h)}.jpg"   # theo tệp nguồn (đổi bảng màu = ảnh mới)
        if not out.exists():
            im = Image.open(src).convert("RGB")
            r_box, r_im = w / h, im.width / im.height
            if r_im > r_box:
                nw = int(im.height * r_box); left = (im.width - nw) // 2
                im = im.crop((left, 0, left + nw, im.height))
            else:
                nh = int(im.width / r_box); top = (im.height - nh) // 2
                im = im.crop((0, top, im.width, top + nh))
            im.thumbnail((int(w * 3), int(h * 3)))
            im.save(out, quality=85)
        return [image(x, y, w, h, out, link, name=name)]

    def add(self, kind, elements, *, dark=False, notes="", chapter=None, footer=True, title="", bg=None):
        self.slides.append(dict(kind=kind, elements=elements, bg=bg or (DARK if dark else "#ffffff"), dark=dark,
                                notes=notes, chapter=chapter, footer=footer, title=title))

    def chrome(self, idx: int, s: dict) -> list:
        """Logo + footer (slogan · ai.slim.vn · số trang). Thêm lúc cuối để biết số trang."""
        els = []
        if not s["footer"]:
            return els
        site = self.b["site"]
        fc = "#8fa0c6" if s["dark"] else MUTED
        if not s["dark"] and s["kind"] != "eco":
            els.append(image(W - M - 26, 22, 26, 26, logo("slimai_icon"), site, fit="contain", name="Logo SlimAI"))
        els += [
            text(M, H - 26, 260, 14, self.b["slogan"], 8, "footer", color="#b9c6e6" if s["dark"] else NAVY, bold=True, name="Slogan"),
            text(W / 2 - 60, H - 26, 120, 14, self.b["site_label"], 8, "footer", color=fc, align="center", link=site, name="ai.slim.vn"),
            text(W - M - 60, H - 26, 60, 14, str(idx + 1), 8, "footer", color=fc, align="right", bold=True, name="Số trang"),
        ]
        return els

    def slogan(self) -> str:
        words = self.b["slogan"].split(" · ")
        return " · ".join(f"[[{c}|{w}]]" for c, w in zip(self.b["slogan_colors"], words))

    def title_block(self, title, subtitle=None):
        els = [text(M, 26, W - 2 * M - 40, 58, title, 26, "title", color=NAVY, bold=True, min_size=20, valign="bottom", name="Tiêu đề")]
        top = 96
        if subtitle:
            els.append(text(M, 88, W - 2 * M, 20, subtitle, 13, "subtitle", color=MUTED, min_size=11, name="Phụ đề"))
            top = 118
        return els, top

    def badge(self, x, y, s, label, size=13):
        return [rect(x, y, s, s, grad=GRAD, radius=s * 0.28, name=f"Huy hiệu {label}"),
                text(x, y, s, s, label, size, "badge", color="#ffffff", bold=True, align="center", valign="middle")]

    # ---------------------------------------------------------- slide tự động
    def cover(self):
        d = self.doc
        ids = d.get("cover_image_ids") or [c["image_id"] for c in d["chapters"] if c.get("image_id")]
        els = []
        if ids:
            els += self.img(ids[0], 400, 0, 320, H)
            els.append(rect(400, 0, 320, H, fill=DARK, alpha=0.35, name="Phủ tối"))
        else:
            els += [rect(470, 60, 200, 200, grad=GRAD, radius=48, alpha=0.9), rect(560, 190, 130, 130, grad=[CYAN, GREEN], radius=32, alpha=0.8)]
        els.append(image(M, 30, 120, 38, logo("slimai_white"), self.b["site"], fit="contain", name="Logo SlimAI"))
        credit = self.b["credit"][d["credit"]]
        els += [
            text(M, 112, 340, 18, d["kicker"], 10, "kicker", color="#9fe3ff", bold=True, caps=True, spacing_chars=1.5),
            text(M, 134, 345, 132, d["title"], 32, "cover-title", color="#ffffff", bold=True, min_size=22, valign="middle"),
            text(M, 272, 340, 50, d["subtitle"], 13, "cover-sub", color="#d9e2f5", min_size=11),
            text(M, 330, 340, 30, [f"{credit} · {self.b['site_label']}", d.get("source_credit") or ""], 10, "cover-credit",
                 color="#ffffff", min_size=8),
        ]
        els.append(text(M, H - 34, 340, 18, self.slogan(), 11, "slogan", color="#8fa0c6", bold=True))
        self.add("cover", els, dark=True, footer=False, title=d["title"], notes=d["subtitle"])

    def agenda(self):
        chs = self.doc["chapters"]
        els, top = self.title_block("Nội dung chính")
        cols = 2 if len(chs) > 4 else 1
        per = -(-len(chs) // cols)
        cw = (W - 2 * M - 20 * (cols - 1)) / cols
        rh = min(72, (H - 50 - top) / per)
        for i, ch in enumerate(chs):
            x = M + (i // per) * (cw + 20)
            y = top + (i % per) * rh
            els += self.badge(x, y + (rh - 6) / 2 - 17, 34, f"{i + 1:02d}", 13)
            els.append(text(x + 42, y, cw - 42, rh - 6, [f"**{ch['title']}**", ch["lead"]], 16, "agenda", color=INK,
                            min_size=11, valign="middle"))
        self.add("agenda", els, title="Nội dung chính", notes="Lộ trình của bài trình bày.")

    def section(self, i, ch):
        els = []
        if ch.get("image_id") and self.images.get(ch["image_id"]):
            els += self.img(ch["image_id"], 360, 0, 360, H)
            els.append(rect(360, 0, 360, H, fill=DARK, alpha=0.25))
        else:
            els += [rect(470, 90, 180, 180, grad=GRAD, radius=44, alpha=0.9)]
        els += self.badge(M, 120, 44, f"{i + 1:02d}", 17)
        els += [
            text(M, 176, 300, 16, f"Chương {i + 1:02d}", 10, "kicker", color="#9fe3ff", bold=True, caps=True, spacing_chars=1.5),
            text(M, 196, 300, 96, ch["title"], 28, "section-title", color="#ffffff", bold=True, min_size=20),
            text(M, 296, 300, 60, ch["lead"], 13, "section-lead", color="#d9e2f5", min_size=11),
        ]
        self.add("section", els, dark=True, chapter=ch["id"], title=ch["title"], notes=ch["lead"])

    def summary(self):
        chs = self.doc["chapters"]
        els, top = self.title_block("Điểm chính cần nhớ")
        n = len(chs)
        cols = 2 if n > 3 else 1
        per = -(-n // cols)
        cw = (W - 2 * M - 16 * (cols - 1)) / cols
        rh = (H - 46 - top) / per
        for i, ch in enumerate(chs):
            x = M + (i // per) * (cw + 16)
            y = top + (i % per) * rh
            els.append(rect(x, y, cw, rh - 8, fill=SOFT, radius=10))
            els += self.badge(x + 10, y + (rh - 8) / 2 - 13, 26, f"{i + 1:02d}", 11)
            els.append(text(x + 46, y + 6, cw - 56, rh - 20, ch["summary"][0] if ch["summary"] else ch["lead"], 14, "summary",
                            min_size=11, valign="middle"))
        self.add("summary", els, title="Điểm chính cần nhớ", notes="Nhắc lại thông điệp chính của từng chương.")

    def course(self):
        c = self.b["course"]
        url = c["url"]
        qr = qr_png(url, self.work / "qr-khoa-hoc.png")
        items = [re.sub(r"<[^>]+>", "", x) for x in c["modules"][0]["items"]][:4]
        items = [x.split(":")[0] if ":" in x[:40] else x for x in items]
        els = [
            text(M, 40, 420, 16, "Bước tiếp theo", 10, "kicker", color="#9fe3ff", bold=True, caps=True, spacing_chars=1.5),
            text(M, 60, 430, 76, c["name"].replace(" cho ", "\ncho ", 1).split("\n"), 28, "section-title", color="#ffffff", bold=True, min_size=20),
            text(M, 140, 430, 40, c["audience"], 12, "section-lead", color="#d9e2f5", min_size=10),
            text(M, 186, 430, 110, items, 14, "course-item", color="#ffffff", bullet=CYAN, space=5, min_size=11),
            rect(M, 306, 190, 32, grad=GRAD, radius=16, name="Nút khoá học"),
            text(M, 306, 190, 32, "Xem khoá học AI Agent →", 12, "button", color="#ffffff", bold=True, align="center",
                 valign="middle", link=url),
            text(M + 200, 314, 230, 16, c["url_label"], 9, "url", color="#9fb0d6", link=url),
            rect(512, 96, 168, 196, fill="#ffffff", radius=14, name="Khung QR"),
            image(526, 110, 140, 140, qr, url, fit="contain", name="QR khoá học"),
            text(512, 256, 168, 26, "QUÉT ĐỂ XEM KHOÁ HỌC", 9, "qr", color=NAVY, bold=True, align="center", valign="middle"),
        ]
        self.add("course", els, dark=True, title=c["name"], notes="Giới thiệu khoá học AI Agent cho người đi làm.")

    def ecosystem(self):
        b = self.b
        site = b["site"]
        els = [image(W / 2 - 90, 24, 180, 58, logo("slimai"), site, fit="contain", name="Logo SlimAI")]
        els.append(text(M, 86, W - 2 * M, 22, self.slogan(), 15, "slogan", color="#8a93a6", bold=True, align="center"))
        cw = (W - 2 * M - 2 * 14) / 3
        for i, e in enumerate(b["ecosystem"]):
            x = M + i * (cw + 14)
            els += [rect(x, 126, cw, 126, fill="#ffffff", radius=12, line=LINE),
                    image(x + 14, 140, cw - 28, 30, logo(e["key"]), e["link"], fit="contain-left", name=e["name"]),
                    text(x + 14, 178, cw - 28, 16, e["label"], 11, "eco-url", color="#0f6fb0", bold=True, link=e["link"]),
                    text(x + 14, 198, cw - 28, 48, e["desc"], 11, "eco-desc", color=INK, min_size=9)]
        els.append(text(M, 266, W - 2 * M, 18, "Theo dõi SlimAI", 12, "follow", color=NAVY, bold=True, align="center"))
        n = len(b["social"])
        gap = 74
        x0 = W / 2 - (n * gap - (gap - 30)) / 2
        for i, s in enumerate(b["social"]):
            x = x0 + i * gap
            els += [image(x, 292, 30, 30, self.social[s["key"]], s["link"], fit="contain", name=s["name"]),
                    text(x - 20, 326, 70, 14, s["name"], 9, "social", color=MUTED, align="center", link=s["link"])]
        self.add("eco", els, title="Hệ sinh thái SlimAI", notes="Kênh và sản phẩm của SlimAI.")

    # ---------------------------------------------------------- slide nội dung (mô hình chọn layout, khung cố định)
    def content(self, s: dict, chapter_id: str):
        lay = s["layout"]
        fn = getattr(self, f"l_{lay}", None)
        if not fn or not fn_ok(s):
            lay = "bullets"
            fn = self.l_bullets
        els, top = self.title_block(s["title"], s.get("subtitle"))
        els += fn(s, top)
        self.add(lay, els, chapter=chapter_id, title=s["title"], notes=s.get("notes") or "")

    def l_bullets(self, s, top):
        items = s.get("bullets") or [s.get("caption") or ""]
        bottom = H - 48
        if s.get("image_id") and self.images.get(s["image_id"]):
            els = [text(M, top + 4, 330, bottom - top - 4, items, 16, "body", bullet=PINK, space=9, min_size=14)]
            els += self.img(s["image_id"], 392, top + 4, W - M - 392, bottom - top - (26 if s.get("caption") else 4))
            if s.get("caption"):
                els.append(text(392, bottom - 18, W - M - 392, 18, s["caption"], 10, "caption", color=MUTED, min_size=9))
            return els
        return [text(M, top + 4, W - 2 * M, bottom - top - 4, items, 18, "body", bullet=PINK, space=12, min_size=14)]

    def l_cards(self, s, top):
        cards = (s.get("cards") or [])[:4]
        n = len(cards)
        gap = 14
        cw = (W - 2 * M - gap * (n - 1)) / n
        ch = H - 52 - top
        els = []
        for i, c in enumerate(cards):
            x = M + i * (cw + gap)
            els += [rect(x, top + 4, cw, ch, fill=SOFT, radius=12),
                    text(x + 14, top + 14, 44, 36, c["icon"], 26, "icon", align="left"),
                    text(x + 14, top + 54, cw - 28, 46, c["title"], 17, "card-title", color=NAVY, bold=True, min_size=13),
                    text(x + 14, top + 104, cw - 28, ch - 112, c["text"], 15, "card-text", min_size=12)]
        return els

    def l_steps(self, s, top):
        steps = (s.get("steps") or [])[:5]
        n = len(steps)
        gap = 16
        cw = (W - 2 * M - gap * (n - 1)) / n
        els = []
        for i, st in enumerate(steps):
            x = M + i * (cw + gap)
            els += self.badge(x, top + 14, 36, str(i + 1), 15)
            els += [text(x, top + 62, cw, 44, st["title"], 16, "step-title", color=NAVY, bold=True, min_size=12),
                    text(x, top + 110, cw, H - 54 - top - 110, st["text"], 14, "step-text", min_size=11)]
        return els

    def l_stats(self, s, top):
        stats = (s.get("stats") or [])[:3]
        n = len(stats)
        gap = 18
        cw = (W - 2 * M - gap * (n - 1)) / n
        els = []
        for i, st in enumerate(stats):
            x = M + i * (cw + gap)
            els += [rect(x, top + 10, cw, 190, fill=SOFT, radius=14),
                    text(x + 16, top + 30, cw - 32, 70, st["value"], 48, "stat-value", color=ACCENTS[i % 5], bold=True, min_size=30),
                    text(x + 16, top + 106, cw - 32, 84, st["label"], 15, "stat-label", color=INK, min_size=11)]
        if s.get("caption"):
            els.append(text(M, top + 212, W - 2 * M, 30, s["caption"], 11, "caption", color=MUTED, min_size=10))
        return els

    def l_compare(self, s, top):
        els = []
        cw = (W - 2 * M - 16) / 2
        ch = H - 52 - top
        for i, (side, fill, acc) in enumerate([(s["left"], "#f3f4f7", MUTED), (s["right"], "#eefaff", CYAN)]):
            x = M + i * (cw + 16)
            els += [rect(x, top + 4, cw, ch, fill=fill, radius=12),
                    text(x + 16, top + 16, cw - 32, 26, side["title"], 18, "compare-title", color=NAVY, bold=True, min_size=13),
                    text(x + 16, top + 48, cw - 32, ch - 56, side["items"], 16, "compare-item", bullet=acc, space=8, min_size=12)]
        return els

    def l_quote(self, s, top):
        return [text(M, top - 10, 70, 110, "“", 80, "quote-mark", color=PINK, bold=True),
                text(M + 70, top + 20, W - 2 * M - 90, 170, s["quote"], 24, "quote", color=NAVY, italic=True, min_size=16, valign="middle"),
                text(M + 70, top + 200, W - 2 * M - 90, 22, f"— {s['by']}" if s.get("by") else "", 13, "quote-by", color=MUTED)]

    def l_prompt(self, s, top):
        ch = 180 if s.get("caption") else 220
        els = [rect(M, top + 6, W - 2 * M, ch, fill="#0f1a33", radius=10, name="Khung prompt"),
               text(M + 16, top + 16, 300, 14, "PROMPT MẪU", 9, "prompt-label", color="#7fd6f5", bold=True, spacing_chars=1),
               text(M + 16, top + 36, W - 2 * M - 32, ch - 44, s["prompt"], 16, "prompt", color="#e6ecff", font=MONO, min_size=11)]
        if s.get("caption"):
            els.append(text(M, top + ch + 16, W - 2 * M, 40, s["caption"], 14, "caption", color=MUTED, min_size=11))
        return els

    def l_image(self, s, top):
        bottom = H - 48
        side = s.get("bullets") or []
        els = self.img(s["image_id"], M, top + 4, 420, bottom - top - 4)
        els.append(text(M + 436, top + 4, W - 2 * M - 436, 70, s.get("caption") or "", 15, "caption-strong", color=NAVY, bold=True, min_size=12))
        if side:
            els.append(text(M + 436, top + 80, W - 2 * M - 436, bottom - top - 80, side, 15, "side", bullet=PINK, space=8, min_size=12))
        return els

    def l_table(self, s, top):
        header, rows = s["header"][:3], [r[:3] for r in (s.get("rows") or [])[:6]]
        cols = len(header)
        cw = (W - 2 * M) / cols
        rh = min(40, (H - 52 - top - 28) / max(1, len(rows)))
        els = [rect(M, top + 4, W - 2 * M, 28, fill=NAVY, radius=6)]
        for j, hd in enumerate(header):
            els.append(text(M + j * cw + 10, top + 4, cw - 20, 28, hd, 14, "th", color="#ffffff", bold=True, valign="middle", min_size=10))
        for i, r in enumerate(rows):
            y = top + 34 + i * rh
            if i % 2:
                els.append(rect(M, y, W - 2 * M, rh, fill="#f8fafd"))
            for j, cell in enumerate(r + [""] * (cols - len(r))):
                els.append(text(M + j * cw + 10, y + 2, cw - 20, rh - 4, cell, 14, "td", valign="middle", min_size=10))
        return els


def fn_ok(s: dict) -> bool:
    lay = s["layout"]
    need = {"hub": ["cards"], "bullets": ["bullets"], "cards": ["cards"], "steps": ["steps"], "stats": ["stats"], "compare": ["left", "right"],
            "quote": ["quote"], "prompt": ["prompt"], "image": ["image_id"], "table": ["header", "rows"]}[lay]
    return all(s.get(k) for k in need)


def build_deck(doc: dict, work: Path, images: dict[str, Path], social: dict[str, Path], style: str = "co-dien") -> Deck:
    if style == "visual":
        from phong_cach_visual import DeckVisual as cls
    elif style == "retron":
        from phong_cach_retron import DeckRetron as cls
    elif style == "saas":
        from phong_cach_saas import DeckSaas as cls
    else:
        cls = Deck
    d = cls(doc, work, images, social)
    d.cover()
    d.agenda()
    for i, ch in enumerate(doc["chapters"]):
        d.section(i, ch)
        for s in ch["slides"]:
            d.content(s, ch["id"])
    d.summary()
    if doc.get("course_enabled", True):
        d.course()
    d.ecosystem()
    for i, s in enumerate(d.slides):
        s["elements"] = s["elements"] + d.chrome(i, s)
        for e in s["elements"]:
            if e["kind"] == "text" and e["font"] == FONT:
                e["font"] = d.sans
    return d


# ---------------------------------------------------------------- HTML backend (đo chữ + in PDF)
TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|\[\[#[0-9a-fA-F]{6}\|.+?\]\])")


def tokens(t: str):
    """Tách chuỗi thành (chữ, đậm?, màu, nhấn?) — **đậm**, *nhấn nghiêng* và [[#hex|chữ màu]]."""
    for part in TOKEN.split(t):
        if not part:
            continue
        if part.startswith("**"):
            yield part[2:-2], True, None, False
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            yield part[1:-1], None, None, True
        elif part.startswith("[["):
            color, txt = part[2:-2].split("|", 1)
            yield txt, None, color, False
        else:
            yield part, None, None, False


def _rich(t: str, e: dict | None = None) -> str:
    out = []
    for txt, bold, color, acc in tokens(t):
        x = html.escape(txt)
        if bold:
            x = f"<b>{x}</b>"
        if acc:
            f = f"font-family:'{e['accent_font']}';font-weight:400;" if e and e.get("accent_font") else ""
            c = f"color:{e['accent']};" if e and e.get("accent") else ""
            n = "font-style:normal;" if e and not e.get("accent_italic", True) else ""
            x = f'<i style="{f}{c}{n}">{x}</i>'
        if color:
            x = f'<span style="color:{color}">{x}</span>'
        out.append(x)
    return "".join(out)


def _css_grad(colors, angle=45):
    return f"linear-gradient({angle + 90}deg,{','.join(colors)})"


def _rgba(hexc, a):
    r, g, b = hex_rgb(hexc)
    return f"rgba({r},{g},{b},{a})"


def shadow_spec(v) -> tuple[str, float]:
    if isinstance(v, str) and "/" in v:
        c, a = v.split("/")
        return c, float(a)
    return "#1e143c", 0.10


def vt_name(e: dict) -> str:
    """Tên morph: phần tử có tên bắt đầu bằng '!!' được ghép cặp giữa các slide (PowerPoint Morph / View Transitions)."""
    n = e.get("name") or ""
    return "vt-" + re.sub(r"[^a-z0-9]+", "-", n[2:].lower()).strip("-") if n.startswith("!!") else ""


def element_html(e: dict, sid: int, eid: int, fit: bool) -> str:
    st = f"left:{e['x']}pt;top:{e['y']}pt;width:{e['w']}pt;height:{e['h']}pt;"
    vt = vt_name(e)
    attr = f' data-i="{eid}"' + (f' data-vt="{vt}"' if vt else "")
    if e["kind"] in ("rect", "oval"):
        if e.get("grad"):
            bg = _css_grad(e["grad"], e.get("grad_angle", 45))
        elif e.get("fill"):
            bg = _rgba(e["fill"], e["alpha"])
        else:
            bg = "transparent"
        border = f"border:{e['line_w']}pt solid {e['line']};" if e["line"] else ""
        op = f"opacity:{e['alpha']};" if e.get("grad") and e["alpha"] < 1 else ""
        radius = "50%" if e["kind"] == "oval" else f"{e['radius']}pt"
        rot = f"transform:rotate({e['rot']:.3f}deg);" if e.get("rot") else ""
        sh = ""
        if e.get("shadow"):   # True = bóng xám nhẹ; "#rrggbb/alpha" = bóng màu (vd panel tím)
            sc, sa = shadow_spec(e["shadow"])
            sh = f"box-shadow:0 {6 if sa <= .12 else 8}pt {18 if sa <= .12 else 22}pt {_rgba(sc, sa)};"
        return f'<div class="el"{attr} style="{st}background:{bg};border-radius:{radius};{border}{op}{rot}{sh}"></div>'
    if e["kind"] == "image":
        fit_css = {"cover": "cover", "contain": "contain", "contain-left": "contain"}[e["fit"]]
        pos = "left center" if e["fit"] == "contain-left" else "center"
        tag = f'<img src="{Path(e["path"]).as_uri()}" style="width:100%;height:100%;object-fit:{fit_css};object-position:{pos}">'
        if e["link"]:
            tag = f'<a href="{html.escape(e["link"])}">{tag}</a>'
        rad = f"border-radius:{e['radius']}pt;overflow:hidden;" if e.get("radius") else ""
        rot = f"transform:rotate({e['rot']:.3f}deg);" if e.get("rot") else ""
        return f'<div class="el"{attr} style="{st}{rad}{rot}">{tag}</div>'
    jc = {"top": "flex-start", "middle": "center", "bottom": "flex-end"}[e["valign"]]
    ps = []
    for p in e["paras"]:
        inner = _rich(p, e)
        if e["link"]:
            inner = f'<a href="{html.escape(e["link"])}" style="color:inherit;text-decoration:none">{inner}</a>'
        if e["bullet"]:
            ind = round(e["size"] * 1.1, 2)
            ps.append(f'<p style="padding-left:{ind}pt;text-indent:-{ind}pt;margin-bottom:{e["space"]}pt">'
                      f'<span style="display:inline-block;width:{ind}pt;text-indent:0;color:{e["bullet"]}">•</span>{inner}</p>')
        else:
            ps.append(f'<p style="margin-bottom:{e["space"]}pt">{inner}</p>')
    style = (f"{st}justify-content:{jc};font-family:'{e['font']}','Apple Color Emoji','Segoe UI Emoji';font-size:{e['size']}pt;color:{e['color']};"
             f"font-weight:{e.get('weight') or (700 if e['bold'] else 400)};font-style:{'italic' if e['italic'] else 'normal'};text-align:{e['align']};"
             f"letter-spacing:{e['cs']}pt;{'text-transform:uppercase;' if e['caps'] else ''}")
    data = f' data-fit="1" data-max="{e["size"]}" data-min="{e["min"]}"' if fit else ""
    return f'<div class="el t" id="e{sid}_{eid}"{data}{attr} style="{style}"><div class="in">{"".join(ps)}</div></div>'


CSS = f"""@page{{size:{W}pt {H}pt;margin:0}}*{{box-sizing:border-box}}html,body{{margin:0;padding:0}}
body{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}
.slide{{position:relative;width:{W}pt;height:{H}pt;overflow:hidden;break-after:page}}
.el{{position:absolute}}.t{{display:flex;flex-direction:column;overflow:hidden;line-height:{LH}}}
.t p{{margin:0;overflow-wrap:break-word}}.t p:last-child{{margin-bottom:0!important}}b{{font-weight:700}}
a{{color:inherit;text-decoration:none}}"""


def slides_html(deck: Deck, fit: bool, img_src=None) -> list:
    out = []
    for sid, s in enumerate(deck.slides):
        els = "".join(element_html(e, sid, eid, fit) for eid, e in enumerate(s["elements"]))
        if img_src:
            els = re.sub(r'src="(file://[^"]+)"', lambda m: f'src="{img_src(m.group(1))}"', els)
        out.append(f'<section class="slide" style="background:{s["bg"]}">{els}</section>')
    return out


def head_links(deck: Deck) -> str:
    return "".join(f'<link rel="stylesheet" href="{u}">' for u in deck.font_links)


def deck_html(deck: Deck, fit: bool) -> str:
    return (f'<!doctype html><html lang="vi"><head><meta charset="utf-8">{head_links(deck)}<style>{CSS}</style></head>'
            f'<body>{"".join(slides_html(deck, fit))}</body></html>')


FIT_JS = """() => {
  const res = {};
  for (const el of document.querySelectorAll('.t[data-fit]')) {
    const inn = el.querySelector('.in');
    let size = parseFloat(el.dataset.max); const min = parseFloat(el.dataset.min);
    const over = () => inn.offsetHeight > el.clientHeight + 0.5 || inn.scrollWidth > el.clientWidth + 0.5;
    while (over() && size > min) { size = Math.max(min, size - 0.5); el.style.fontSize = size + 'pt';
      inn.querySelectorAll('p').forEach(p => { if (p.style.paddingLeft) { const ind = (size*1.1).toFixed(2);
        p.style.paddingLeft = ind+'pt'; p.style.textIndent = '-'+ind+'pt'; p.firstChild.style.width = ind+'pt'; } }); }
    res[el.id] = {size, over: over()};
  }
  return res;
}"""


def _load(deck: Deck, page, fit: bool) -> None:
    """Ghi HTML ra tệp rồi mở bằng file:// để Chromium đọc được ảnh trên máy."""
    f = deck.work / ("_slides_fit.html" if fit else "_slides.html")
    f.write_text(deck_html(deck, fit=fit))
    page.goto(f.as_uri(), wait_until="networkidle")


def fit_text(deck: Deck, page) -> list[dict]:
    """Đo trong Chromium, co chữ tới khi vừa, rồi đồng bộ cỡ chữ theo cấp (role) trên từng slide. Trả về danh sách slide tràn."""
    _load(deck, page, fit=True)
    page.evaluate("document.fonts.ready")
    res = page.evaluate(FIT_JS)
    problems = {}
    for sid, s in enumerate(deck.slides):
        by_role: dict[str, float] = {}
        for eid, e in enumerate(s["elements"]):
            r = res.get(f"e{sid}_{eid}")
            if e["kind"] != "text" or not r:
                continue
            e["size"] = r["size"]
            by_role[e["role"]] = min(by_role.get(e["role"], 999), r["size"])
            if r["over"]:
                problems.setdefault(sid, []).append(f"ô '{e['role']}' ({' / '.join(e['paras'])[:60]}…)")
        for e in s["elements"]:
            if e["kind"] == "text" and e["role"] in by_role:
                e["size"] = by_role[e["role"]]
    return [{"slide": sid, "problems": p} for sid, p in problems.items()]


def write_pdf(deck: Deck, page, out: Path) -> None:
    _load(deck, page, fit=False)
    page.evaluate("document.fonts.ready")
    page.pdf(path=str(out), width=f"{W / 72}in", height=f"{H / 72}in", print_background=True, prefer_css_page_size=True)


# ---------------------------------------------------------------- PPTX backend (đối tượng gốc, sửa được)
def write_pptx(deck: Deck, out: Path, title: str) -> None:
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(W), Pt(H)
    prs.core_properties.title = title
    prs.core_properties.author = "SlimAI (ai.slim.vn)"
    blank = prs.slide_layouts[6]
    rgb = lambda c: RGBColor(*hex_rgb(c))

    def set_alpha(fill_fore_xml_parent, alpha):
        if alpha >= 1:
            return
        for clr in fill_fore_xml_parent.iter(qn("a:srgbClr")):
            a = etree.SubElement(clr, qn("a:alpha"))
            a.set("val", str(int(alpha * 100000)))

    def grad_fill(shape, colors, angle=45, alpha=1.0):
        spPr = shape._element.spPr
        for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
            for x in spPr.findall(qn(tag)):
                spPr.remove(x)
        g = etree.Element(qn("a:gradFill"), rotWithShape="1")
        gs = etree.SubElement(g, qn("a:gsLst"))
        for i, c in enumerate(colors):
            s = etree.SubElement(gs, qn("a:gs"), pos=str(int(i * 100000 / (len(colors) - 1))))
            clr = etree.SubElement(s, qn("a:srgbClr"), val=c.lstrip("#").upper())
            if alpha < 1:
                etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
        etree.SubElement(g, qn("a:lin"), ang=str(int(angle * 60000)), scaled="0")
        geom = spPr.find(qn("a:prstGeom"))
        geom.addnext(g)

    def shadow(shp, spec=True):
        sc, sa = shadow_spec(spec)
        big = sa > .12
        spPr = shp._element.spPr
        eff = etree.SubElement(spPr, qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(Pt(22 if big else 18)), dist=str(Pt(8 if big else 6)),
                              dir="5400000", algn="t", rotWithShape="0")
        c = etree.SubElement(sh, qn("a:srgbClr"), val=sc.lstrip("#").upper())
        etree.SubElement(c, qn("a:alpha"), val=str(int(sa * 100000)))

    def keep_link_color(run):
        """Chữ có link giữ màu riêng (PowerPoint 365) thay vì đổi sang màu hyperlink xanh của theme."""
        h = run._r.find(qn("a:rPr")).find(qn("a:hlinkClick"))
        if h is not None:
            ext = etree.SubElement(etree.SubElement(h, qn("a:extLst")), qn("a:ext"), uri="{A12FA001-AC4F-418D-AE19-62706E023703}")
            etree.SubElement(ext, "{http://schemas.microsoft.com/office/drawing/2018/hyperlinkcolor}hlinkClr", val="tx")

    MORPH = (
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" '
        'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" Requires="p159">'
        '<p:transition spd="slow" p14:dur="1100"><p159:morph option="byObject"/></p:transition></mc:Choice>'
        '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')

    for s in deck.slides:
        sl = prs.slides.add_slide(blank)
        if deck.transition == "morph":
            sl._element.find(qn("p:clrMapOvr")).addnext(etree.fromstring(MORPH))
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = rgb(s["bg"])
        for e in s["elements"]:
            x, y, w, h = Pt(e["x"]), Pt(e["y"]), Pt(e["w"]), Pt(e["h"])
            if e["kind"] in ("rect", "oval"):
                kind = MSO_SHAPE.OVAL if e["kind"] == "oval" else (MSO_SHAPE.ROUNDED_RECTANGLE if e["radius"] else MSO_SHAPE.RECTANGLE)
                shp = sl.shapes.add_shape(kind, x, y, w, h)
                if e["kind"] == "rect" and e["radius"]:
                    shp.adjustments[0] = min(0.5, e["radius"] / min(e["w"], e["h"]))
                if e.get("grad"):
                    grad_fill(shp, e["grad"], angle=e.get("grad_angle", 45), alpha=e["alpha"])
                elif e.get("fill"):
                    shp.fill.solid()
                    shp.fill.fore_color.rgb = rgb(e["fill"])
                    set_alpha(shp._element.spPr, e["alpha"])
                else:
                    shp.fill.background()
                if e["line"]:
                    shp.line.color.rgb = rgb(e["line"])
                    shp.line.width = Pt(e.get("line_w", 0.75))
                else:
                    shp.line.fill.background()
                shp.shadow.inherit = False
                if e.get("shadow"):
                    shadow(shp, e["shadow"])
                if e.get("rot"):
                    shp.rotation = e["rot"]
                shp.text_frame.text = ""
                if e["name"]:
                    shp.name = e["name"]
            elif e["kind"] == "image":
                src = Path(e["path"])
                if e["fit"] == "cover":
                    pic = sl.shapes.add_picture(str(src), x, y, w, h)
                else:  # contain: giữ tỉ lệ, căn giữa hoặc căn trái trong khung
                    with Image.open(src) as im:
                        r = im.width / im.height
                    bw, bh = e["w"], e["h"]
                    pw, ph = (bw, bw / r) if bw / bh < r else (bh * r, bh)
                    px = e["x"] if e["fit"] == "contain-left" else e["x"] + (bw - pw) / 2
                    pic = sl.shapes.add_picture(str(src), Pt(px), Pt(e["y"] + (bh - ph) / 2), Pt(pw), Pt(ph))
                if e.get("radius"):   # ảnh bo góc: đổi hình học ảnh sang roundRect
                    geom = pic._element.spPr.find(qn("a:prstGeom"))
                    geom.set("prst", "roundRect")
                    av = geom.find(qn("a:avLst"))
                    if av is None:
                        av = etree.SubElement(geom, qn("a:avLst"))
                    gd = etree.SubElement(av, qn("a:gd"), name="adj")
                    gd.set("fmla", f"val {int(min(50000, e['radius'] / min(e['w'], e['h']) * 100000))}")
                if e.get("rot"):
                    pic.rotation = e["rot"]
                if e["link"]:
                    pic.click_action.hyperlink.address = e["link"]
                if e["name"]:
                    pic.name = e["name"]
            else:
                tb = sl.shapes.add_textbox(x, y, w, h)
                if e["name"]:
                    tb.name = e["name"]
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[e["valign"]]
                for i, para in enumerate(e["paras"]):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[e["align"]]
                    p.line_spacing = Pt(e["size"] * LH)
                    if i < len(e["paras"]) - 1:
                        p.space_after = Pt(e["space"])
                    if e["bullet"]:
                        ind = int(Pt(e["size"] * 1.1))
                        pPr = p._p.get_or_add_pPr()
                        pPr.set("marL", str(ind)); pPr.set("indent", str(-ind))
                        bc = etree.SubElement(pPr, qn("a:buClr"))
                        etree.SubElement(bc, qn("a:srgbClr"), val=e["bullet"].lstrip("#").upper())
                        etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
                        etree.SubElement(pPr, qn("a:buChar"), char="•")
                    txt = para.upper() if e["caps"] else para
                    for part, bold, color, acc in tokens(txt):
                        r = p.add_run()
                        r.text = part
                        f = r.font
                        f.name = (e.get("accent_font") or e["font"]) if acc else e["font"]
                        f.size = Pt(e["size"])
                        f.bold = bool(e["bold"] or bold) and not (acc and e.get("accent_font"))
                        f.italic = bool(e["italic"] or (acc and e.get("accent_italic", True)))
                        f.color.rgb = rgb(color or (e.get("accent") if acc else None) or e["color"])
                        if e["cs"]:
                            r._r.get_or_add_rPr().set("spc", str(int(e["cs"] * 100)))
                        if e["link"]:
                            r.hyperlink.address = e["link"]
                            keep_link_color(r)
        if s["notes"]:
            sl.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(str(out))


def render(doc: dict, work: Path, images: dict[str, Path], out_pptx: Path, out_pdf: Path, page, shorten=None, rounds: int = 2,
           style: str = "co-dien", out_html: Path | None = None) -> dict:
    """Dựng bộ slide. `shorten(slide_json, problems) -> slide_json` dùng để rút gọn slide tràn chữ (tối đa `rounds` vòng)."""
    social = social_pngs(work / "icon", page)
    doc = copy.deepcopy(doc)
    for r in range(rounds + 1):
        deck = build_deck(doc, work, images, social, style)
        overflow = fit_text(deck, page)
        if not overflow or not shorten or r == rounds:
            break
        fixed = 0
        for o in overflow:
            s = deck.slides[o["slide"]]
            for ch in doc["chapters"]:
                for k, src in enumerate(ch["slides"]):
                    if src["title"] == s["title"] and ch["id"] == s["chapter"]:
                        print(f"  · rút gọn slide {o['slide'] + 1}: {s['title'][:50]}", flush=True)
                        try:
                            ch["slides"][k] = shorten(src, o["problems"])
                            fixed += 1
                        except RuntimeError as exc:
                            print(f"  ⚠ không rút gọn được ({exc}); giữ nguyên và ghi vào báo cáo", flush=True)
        if not fixed:
            break
    write_pdf(deck, page, out_pdf)
    write_pptx(deck, out_pptx, doc["title"])
    if out_html:
        from trinh_chieu import write_presenter
        write_presenter(deck, out_html, doc["title"])
    return {"doc": doc, "deck": deck, "overflow": overflow}
