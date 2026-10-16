"""Tác tử đọc nguồn: PDF, DOCX, PPTX, MD/TXT, HTML/URL -> Markdown có mốc nguồn + danh mục ảnh minh hoạ."""
from __future__ import annotations

import hashlib
import io
import json
import re
from pathlib import Path

from PIL import Image

MIN_IMAGE_PX = 300          # bỏ logo, icon, ảnh trang trí nhỏ
MAX_IMAGES = 60             # đủ cho biên tập chọn, không làm prompt quá dài
SUPPORTED = {".pdf", ".docx", ".pptx", ".md", ".markdown", ".txt", ".html", ".htm"}


class SourceError(RuntimeError):
    pass


def _save_image(data: bytes, images_dir: Path, seen: set[str], meta: dict, catalog: list[dict]) -> None:
    digest = hashlib.sha1(data).hexdigest()[:12]
    if digest in seen:
        return
    try:
        im = Image.open(io.BytesIO(data))
        im.load()
    except Exception:
        return
    if min(im.size) < MIN_IMAGE_PX:
        return
    if im.mode not in ("RGB", "L"):
        bg = Image.new("RGB", im.size, (255, 255, 255))
        im = im.convert("RGBA")
        bg.paste(im, mask=im.split()[3])
        im = bg
    im = im.convert("RGB")
    im.thumbnail((1600, 1600))
    seen.add(digest)
    image_id = f"anh{len(catalog) + 1:02d}"
    path = images_dir / f"{image_id}.jpg"
    im.save(path, quality=82, optimize=True)
    catalog.append({"id": image_id, "file": path.name, "w": im.width, "h": im.height, **meta})


def _context(text: str, limit: int = 220) -> str:
    return re.sub(r"\s+", " ", text).strip()[:limit]


def read_pdf(path: Path, images_dir: Path, seen: set, catalog: list) -> str:
    import fitz

    doc = fitz.open(path)
    pages = [page.get_text("text") for page in doc]
    # Dòng lặp trên ≥ 40% số trang (header, footer, số trang) là nhiễu -> bỏ
    from collections import Counter
    freq = Counter(l.strip() for t in pages for l in set(t.splitlines()) if l.strip())
    noise = {l for l, c in freq.items() if len(pages) >= 5 and c >= 0.4 * len(pages)}
    parts = []
    for i, page in enumerate(doc, 1):
        text = "\n".join(l for l in pages[i - 1].splitlines() if l.strip() not in noise and not re.fullmatch(r"\s*\d{1,3}\s*", l))
        parts.append(f"\n[nguồn: {path.name}, trang {i}]\n{text.strip()}")
        for info in page.get_images(full=True):
            xref = info[0]
            try:
                pix = fitz.Pixmap(doc, xref)
                if pix.alpha or pix.n > 3:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                data = pix.tobytes("png")
            except Exception:
                continue
            _save_image(data, images_dir, seen, {"source": path.name, "page": i, "context": _context(text)}, catalog)
    return "\n".join(parts)


def read_docx(path: Path, images_dir: Path, seen: set, catalog: list) -> str:
    import docx

    d = docx.Document(str(path))
    lines = [f"[nguồn: {path.name}]"]
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower()
        if style.startswith("heading"):
            level = re.sub(r"\D", "", style) or "2"
            lines.append("#" * min(int(level) + 1, 4) + " " + t)
        elif "list" in style:
            lines.append("- " + t)
        else:
            lines.append(t)
    for table in d.tables:
        rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in table.rows]
        if rows:
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("|" + "---|" * len(rows[0]))
            lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    body = "\n".join(lines)
    for rel in d.part.rels.values():
        if "image" in rel.reltype:
            _save_image(rel.target_part.blob, images_dir, seen, {"source": path.name, "page": None, "context": ""}, catalog)
    return body


def read_pptx(path: Path, images_dir: Path, seen: set, catalog: list) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                _save_image(shape.image.blob, images_dir, seen,
                            {"source": path.name, "page": i, "context": _context(" ".join(texts))}, catalog)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        parts.append(f"\n[nguồn: {path.name}, slide {i}]\n" + "\n".join(texts) + (f"\nGhi chú: {notes}" if notes else ""))
    return "\n".join(parts)


def read_html(text: str, name: str) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "form", "noscript"]):
        tag.decompose()
    lines = [f"[nguồn: {name}]"]
    for el in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "td"]):
        t = el.get_text(" ", strip=True)
        if not t:
            continue
        if el.name in ("h1", "h2", "h3", "h4"):
            lines.append("#" * int(el.name[1]) + " " + t)
        elif el.name == "li":
            lines.append("- " + t)
        else:
            lines.append(t)
    return "\n".join(lines)


def read_url(url: str) -> str:
    import requests

    resp = requests.get(url, timeout=30, headers={"User-Agent": "Mozilla/5.0 slide-ebook-agent"})
    resp.raise_for_status()
    return read_html(resp.text, url)


def read_sources(sources: list[str], work_dir: Path) -> tuple[str, list[dict]]:
    """Trả về (markdown gộp, danh mục ảnh). Ảnh lưu ở work_dir/anh/."""
    images_dir = work_dir / "anh"
    images_dir.mkdir(parents=True, exist_ok=True)
    seen: set[str] = set()
    catalog: list[dict] = []
    parts = []
    for src in sources:
        if re.match(r"https?://", src):
            parts.append(read_url(src))
            continue
        path = Path(src).expanduser()
        if not path.exists():
            raise SourceError(f"Không tìm thấy tệp nguồn: {path}")
        ext = path.suffix.lower()
        if ext not in SUPPORTED:
            raise SourceError(f"Chưa hỗ trợ định dạng {ext} ({path.name}). Hỗ trợ: {', '.join(sorted(SUPPORTED))}")
        if ext == ".pdf":
            parts.append(read_pdf(path, images_dir, seen, catalog))
        elif ext == ".docx":
            parts.append(read_docx(path, images_dir, seen, catalog))
        elif ext == ".pptx":
            parts.append(read_pptx(path, images_dir, seen, catalog))
        elif ext in (".html", ".htm"):
            parts.append(read_html(path.read_text(errors="ignore"), path.name))
        else:
            parts.append(f"[nguồn: {path.name}]\n" + path.read_text(errors="ignore"))
    if len(catalog) > MAX_IMAGES:   # rải đều trên toàn tài liệu thay vì lấy các trang đầu
        step = len(catalog) / MAX_IMAGES
        keep = {catalog[int(i * step)]["id"] for i in range(MAX_IMAGES)}
        for c in catalog:
            if c["id"] not in keep:
                (images_dir / c["file"]).unlink(missing_ok=True)
        catalog[:] = [c for c in catalog if c["id"] in keep]
    markdown = "\n\n".join(parts).strip()
    if len(re.sub(r"\s", "", markdown)) < 200:
        raise SourceError("Nguồn gần như không có chữ (có thể là PDF scan). Agent không đọc OCR; hãy gửi bản có chữ hoặc transcript.")
    (work_dir / "nguon.md").write_text(markdown)
    (work_dir / "anh.json").write_text(json.dumps(catalog, ensure_ascii=False, indent=1))
    return markdown, catalog
